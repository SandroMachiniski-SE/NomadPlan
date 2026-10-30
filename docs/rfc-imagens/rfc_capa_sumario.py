"""Atualiza a capa e a identificação do RFC.docx e troca o sumário digitado à mão por um sumário automático.

Uso: depois de rfc_secoes_8_13.py; em seguida rfc_para_pdf.py, que calcula as páginas do sumário no Word
antes de gerar o PDF. Pode rodar de novo sobre o RFC já atualizado."""
import re
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

RFC = Path(__file__).resolve().parents[2] / "RFC.docx"
DATA = "novembro de 2026"
VERSAO = "2.0"
# Números de seção seguidos do título ("6", "6.3"); só os níveis 1 e 2 entram no sumário.
TITULO = re.compile(r"^(\d+(?:\.\d+)*)\.?\s+\S")

doc = Document(RFC)
ps = doc.paragraphs


def trocar_texto(p, novo):
    """Troca o texto mantendo a formatação do primeiro trecho (negrito, fonte, tamanho)."""
    p.runs[0].text = novo
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


def eh_titulo(p):
    return bool(p.runs) and bool(p.runs[0].bold) and len(p.text.strip()) < 110


def nivel_de_estrutura(p, nivel):
    ppr = p._p.get_or_add_pPr()
    for antigo in ppr.findall(qn("w:outlineLvl")):
        ppr.remove(antigo)
    el = OxmlElement("w:outlineLvl")
    el.set(qn("w:val"), str(nivel))
    ppr.insert_element_before(el, "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")


def campo_sumario():
    """Parágrafo com o campo TOC: títulos de nível 1 e 2 (pelo nível de estrutura), com links."""
    p = OxmlElement("w:p")
    for tipo, texto in (("begin", None), (None, ' TOC \\o "1-2" \\h \\z \\u '), ("separate", None),
                        (None, "Atualize o sumário no Word: botão direito > Atualizar campo."), ("end", None)):
        r = OxmlElement("w:r")
        if tipo:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), tipo)
            if tipo == "begin":
                fc.set(qn("w:dirty"), "true")
            r.append(fc)
        else:
            t = OxmlElement("w:instrText" if "TOC" in texto else "w:t")
            t.set(qn("xml:space"), "preserve")
            t.text = texto
            r.append(t)
        p.append(r)
    return p


# ---- capa e identificação
for p in ps[:12]:
    if re.match(r"Sandro Marcos Machiniski\s*-", p.text.strip()):
        trocar_texto(p, f"Sandro Marcos Machiniski - {DATA}")

i_sumario = next(i for i, p in enumerate(ps) if p.text.strip() == "Sumário")
i_ident = next(i for i, p in enumerate(ps) if i > i_sumario and p.text.strip() == "Identificação" and eh_titulo(p))

for p in ps[i_ident:i_ident + 10]:
    t = p.text.strip()
    if t.startswith("Data:"):
        trocar_texto(p, f"Data: {DATA[0].upper()}{DATA[1:]}")
    elif t.startswith("Versão:"):
        trocar_texto(p, f"Versão: {VERSAO}")

# ---- o modelo numerava duas subseções como 3.3
for p in ps:
    if p.text.strip() == "3.3 Análise de Soluções Existentes (Benchmark)":
        trocar_texto(p, "3.4 Análise de Soluções Existentes (Benchmark)")

# ---- sumário: tudo entre o título "Sumário" e "Identificação" vira o campo automático
titulo_sumario = ps[i_sumario]._p
identificacao = ps[i_ident]._p
elemento = titulo_sumario.getnext()
while elemento is not identificacao:
    proximo = elemento.getnext()
    elemento.getparent().remove(elemento)
    elemento = proximo
titulo_sumario.addnext(campo_sumario())

# capa sozinha na primeira página; sumário e identificação começam páginas novas
ps = doc.paragraphs
for p in (ps[i_sumario], next(p for p in ps if p.text.strip() == "Identificação" and eh_titulo(p))):
    p.paragraph_format.page_break_before = True

# ---- níveis de estrutura dos títulos (o sumário e os marcadores do PDF usam esses níveis)
titulos = 0
depois_do_sumario = False
for p in doc.paragraphs:
    t = p.text.strip()
    if t == "Identificação" and eh_titulo(p):
        depois_do_sumario = True
        nivel_de_estrutura(p, 0)
        titulos += 1
        continue
    if not depois_do_sumario or not eh_titulo(p):
        continue
    m = TITULO.match(t)
    if m and m.group(1).count(".") <= 1:
        nivel_de_estrutura(p, m.group(1).count("."))
        titulos += 1

doc.save(RFC)
print("ok:", titulos, "títulos no sumário")
